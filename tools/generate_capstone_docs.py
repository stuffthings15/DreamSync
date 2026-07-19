from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
NAVY, BLUE, TEAL, LIGHT, GRAY, RED = "102A43", "167DAD", "1F9D8A", "EAF3F8", "52606D", "B42318"

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd")) or OxmlElement("w:shd")
    if shd.getparent() is None: tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def margins(cell, value=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar") or OxmlElement("w:tcMar")
    if tc_mar.getparent() is None: tc_pr.append(tc_mar)
    for name in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{name}")) or OxmlElement(f"w:{name}")
        if node.getparent() is None: tc_mar.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")

def usable_twips(doc):
    s = doc.sections[0]
    return round((s.page_width - s.left_margin - s.right_margin) / 635)

def fixed_table(table, widths, indent=90):
    table.autofit = False
    total = sum(widths)
    pr = table._tbl.tblPr
    tbl_w = pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW"); pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa"); tbl_w.set(qn("w:w"), str(total))
    tbl_ind = pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd"); pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa"); tbl_ind.set(qn("w:w"), str(indent))
    layout = pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout"); pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    for grid_col, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW"); cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:type"), "dxa"); tc_w.set(qn("w:w"), str(width))

def setup(doc, name, compact=False):
    s = doc.sections[0]
    edge = .48 if compact else .65
    s.top_margin = s.bottom_margin = Inches(edge)
    s.left_margin = s.right_margin = Inches(.58 if compact else .7)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"; normal.font.size = Pt(8.2 if compact else 9.3)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(2.3 if compact else 4)
    title = doc.styles["Title"]
    title.font.name = "Aptos Display"; title.font.size = Pt(21 if compact else 24)
    title.font.bold = True; title.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.space_after = Pt(3)
    for style, size in (("Heading 1", 14), ("Heading 2", 10.5)):
        st = doc.styles[style]
        st.font.name = "Aptos Display"; st.font.size = Pt(size); st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(BLUE)
        st.paragraph_format.space_before = Pt(6); st.paragraph_format.space_after = Pt(2)
    h = s.header.paragraphs[0]
    h.text = "DREAM SYNC  /  PRIVATE CAPSTONE RESEARCH PROTOTYPE"; h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in h.runs: r.font.size = Pt(7); r.font.bold = True; r.font.color.rgb = RGBColor.from_string(BLUE)
    f = s.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = f.add_run("Experimental • Non-medical • Human review required • 18 July 2026")
    r.font.size = Pt(7); r.font.color.rgb = RGBColor.from_string(GRAY)
    doc.core_properties.title = name; doc.core_properties.author = "Dream Sync project"

def title(doc, name, sub):
    doc.add_paragraph(name, style="Title")
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(5)
    r = p.add_run(sub); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string(TEAL)

def heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p._p.get_or_add_pPr().append(OxmlElement("w:keepNext"))

def bullets(parent, items, size=None):
    for item in items:
        p = parent.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(.14); p.paragraph_format.first_line_indent = Inches(-.1)
        p.paragraph_format.space_after = Pt(1.2)
        r = p.add_run(item)
        if size: r.font.size = Pt(size)

def boxes(doc, blocks):
    t = doc.add_table(rows=1, cols=len(blocks)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    total = usable_twips(doc) - 85
    widths = [total // len(blocks)] * len(blocks)
    widths[-1] += total - sum(widths)
    fixed_table(t, widths, 85)
    for i, (head, body) in enumerate(blocks):
        c = t.cell(0, i); shade(c, LIGHT); margins(c, 85)
        r = c.paragraphs[0].add_run(head.upper()); r.bold = True; r.font.size = Pt(7.7); r.font.color.rgb = RGBColor.from_string(BLUE)
        if isinstance(body, list): bullets(c, body, 7.45)
        else:
            p = c.add_paragraph(body); p.paragraph_format.space_after = Pt(1)
            for r in p.runs: r.font.size = Pt(7.55)

def kv(doc, rows):
    t = doc.add_table(rows=0, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (key, value) in enumerate(rows):
        a, b = t.add_row().cells; shade(a, BLUE if i % 2 == 0 else NAVY); shade(b, LIGHT if i % 2 == 0 else "F5F8FA")
        for c in (a, b): margins(c); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        r = a.paragraphs[0].add_run(key); r.bold = True; r.font.color.rgb = RGBColor(255,255,255)
        b.paragraphs[0].add_run(value)
    total = usable_twips(doc) - 90
    fixed_table(t, [round(total * .22), total - round(total * .22)])
    return t

def proposal():
    d = Document(); setup(d, "Dream Sync One-Page Proposal", True)
    title(d, "Dream Sync", "A private, safety-gated phone + watch experiment for conservative overnight audio cueing")
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    r = p.add_run("VALUE  "); r.bold = True; r.font.color.rgb = RGBColor.from_string(BLUE)
    p.add_run("Turn a lab-inspired cue concept into a reproducible proof that logs what happened, fails safely, and makes no sleep-effect claim before evidence exists.")
    boxes(d, [
        ("Problem", "Consumer wearables lack EEG timing. Naive cueing can wake the user, mislabel sleep, or create claims the data cannot support."),
        ("User", "Nathan using an S23+ and Galaxy Watch Ultra at home. Today there is no synchronized, inspectable private protocol."),
        ("Riskiest assumption", "Motion + heart-rate quality gates can find conservative cue-eligible windows without unsafe repeat playback."),
    ])
    heading(d, "Workflow", 2)
    kv(d, [
        ("INPUT", "Motion + heart rate, calibrated phone audio, configuration, blinded ACTIVE/SHAM assignment."),
        ("LOCAL", "Quality checks → estimated sleep onset → NREM-eligible / REM-like eligible → cue policy."),
        ("DELIVERY", "One bound pair; expiring cue; requested → accepted → audio started → completed/failed."),
        ("REVIEW", "Survey + export → de-identified Codex summary → source/limits check → human approval."),
    ])
    heading(d, "Success metrics", 2)
    boxes(d, [
        ("Reliability", ["3/3 accelerated device runs recover a summary.", "One valid cue: at most one playback. Invalid: none."]),
        ("Research readiness", ["Plan 12–20 usable sensor-only, ACTIVE, and SHAM nights.", "Every usable night has quality, lifecycle, terminal state, survey."]),
        ("Claim gate", ["Baseline not measured yet.", "No sleep-effect claim before trial + human review."]),
    ])
    heading(d, "Scope fence", 2)
    boxes(d, [
        ("Must", ["Private APKs; readiness/start/stop; Room; blind trial; export/replay; tests.", "Visible failure reasons; non-medical language."]),
        ("Won't", ["Play launch, production signing, Samsung SDK, EEG phase locking, cloud.", "Guaranteed lucid dreams or medical claims."]),
        ("Risk → control", ["Wrong stage → eligibility labels.", "Duplicate/stale → fail closed.", "Audio/sensor loss → logging only + reason."]),
    ])
    heading(d, "Three-week finish line", 2)
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(2)
    for head, body in (("W1 ", "protocol, persistence, unit tests  |  "), ("W2 ", "audio/quality gates, sham, survey, export  |  "), ("W3 ", "hardware runs, replay, report, capstone package")):
        p.add_run(head).bold = True; p.add_run(body)
    final_table = d.add_table(rows=1, cols=1); fixed_table(final_table, [usable_twips(d) - 70], 70)
    c = final_table.cell(0,0); shade(c, NAVY); margins(c, 70)
    r = c.paragraphs[0].add_run("RESEARCH BASIS  Besedovsky et al. (2017), DOI 10.1038/s41467-017-02170-3: paired 50 ms pink-noise tones, 1.075 s apart, EEG phase-locked in a 14-man lab crossover. Dream Sync preserves timing only; it does not reproduce EEG phase locking.")
    r.font.size = Pt(7.1); r.font.color.rgb = RGBColor(255,255,255)
    d.save(DOCS / "Dream Sync One-Page Proposal.docx")

REQS = [
    ("Protocol identity", "Version, session/message/cue IDs, sender, target, UTC, TTL, payload.", "shared protocol", "Round-trip + hostile payload", "PASS"),
    ("Device binding", "One capability-selected phone/watch; direct messages only.", "MainActivity; SessionConfig; services", "Hardware pending", "IMPLEMENTED"),
    ("Durable transport", "Data Items config/state/summary; Message cue/ACK; Asset CSV.", "listeners + services", "Build + source map", "IMPLEMENTED"),
    ("At-most-once", "Room claim suppresses duplicate cue playback.", "CuePlaybackService; PhoneDao", "Property tests; hardware pending", "IMPLEMENTED"),
    ("Cue lifecycle", "Allowance only after AUDIO_STARTED.", "phone/wear event flow", "Build + policy tests", "IMPLEMENTED"),
    ("Persistence", "Phone/watch sessions, events, survey, terminal reason.", "Room databases + schemas", "Schema generation", "PASS"),
    ("Process death", "Old running session interrupted; fresh handshake required.", "WearDao; service", "Instrumentation pending", "IMPLEMENTED"),
    ("Watchdogs", "Phone, streams, storage, permission, battery gates.", "SleepSessionService", "Hardware pending", "IMPLEMENTED"),
    ("Audio readiness", "Preload, speaker route, volume, focus, start/completion.", "AudioOutput", "Timing unit; routes pending", "IMPLEMENTED"),
    ("Signal quality", "Freshness/coverage; lucid requires heart.", "quality monitor + service", "SafetyCoreTest", "PASS"),
    ("Sleep onset", "Windows anchor to sustained quiet onset.", "SleepOnsetEstimator", "SafetyCoreTest", "PASS"),
    ("Language", "NREM-eligible / REM-like eligible; non-medical.", "engine + UI + docs", "Text inspection", "PASS"),
    ("Blinding", "ACTIVE/SHAM hidden until survey.", "MainActivity; config", "Live UI pending", "IMPLEMENTED"),
    ("Morning workflow", "Ratings, awakenings, awareness, recall, lucid, adverse effects, notes.", "survey entity + UI", "UI test pending", "IMPLEMENTED"),
    ("Data control", "History, versioned export, deletion.", "MainActivity + DAOs", "File/UI test pending", "IMPLEMENTED"),
    ("Replay", "Same input/version → same decisions/fingerprint.", "ReplayHarness", "ReplayHarnessTest", "PASS"),
    ("AI gate", "De-identified summary, sources, limits, failures, human review.", "Codex report", "Document inspection", "PASS"),
    ("Private package", "Phone/watch debug APKs; no public release.", "Gradle outputs; README", "assembleDebug", "PASS"),
    ("Hardware", "S23+/Watch Ultra active/sham/sensor/fault/battery/recovery.", "Verification Results", "Physical devices required", "PENDING"),
]

def matrix():
    d = Document(); setup(d, "Dream Sync Requirement-Evidence Matrix")
    title(d, "Requirement–Evidence Matrix", "Traceability from plan to code, tests, and remaining proof")
    p = d.add_paragraph("PASS = locally verified. IMPLEMENTED = built, physical/UI validation open. PENDING = not honestly testable here."); p.runs[0].italic = True
    t = d.add_table(rows=1, cols=5); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fixed_table(t, [1550, 2900, 2250, 2350, usable_twips(d) - 9140])
    heads = ("Requirement", "Acceptance", "Evidence", "Verification", "Status")
    for i,h in enumerate(heads):
        c=t.rows[0].cells[i]; shade(c,NAVY); margins(c); r=c.paragraphs[0].add_run(h); r.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for ri,row in enumerate(REQS):
        cells=t.add_row().cells
        for i,value in enumerate(row):
            shade(cells[i], "F5F8FA" if ri%2 else "FFFFFF"); margins(cells[i],90)
            r=cells[i].paragraphs[0].add_run(value); r.font.size=Pt(7.3)
            if i==0: r.bold=True
            if i==4: r.bold=True; r.font.color.rgb=RGBColor.from_string(TEAL if value=="PASS" else BLUE if value=="IMPLEMENTED" else RED)
    fixed_table(t, [1550, 2900, 2250, 2350, usable_twips(d) - 9140])
    heading(d,"Acceptance boundary")
    bullets(d,["Software tests do not prove sleep effectiveness.","Hardware acceptance stays open until real logs and recovered summaries exist.","No sleep-effect claim before the 12–20-night trial.","Public release stays blocked by license, signing, policy, and release work."])
    d.save(DOCS / "Dream Sync Requirement-Evidence Matrix.docx")

def report():
    d=Document(); setup(d,"Dream Sync Codex Research Report")
    title(d,"Codex Research Report","Validated, de-identified capstone synthesis — DRAFT FOR HUMAN REVIEW")
    p=d.add_paragraph(); r=p.add_run("REVIEW GATE: Do not use for a sleep-effect claim until a human verifies the source mapping and the device study is complete."); r.bold=True; r.font.color.rgb=RGBColor.from_string(RED)
    heading(d,"Executive finding")
    d.add_paragraph("Dream Sync is structured as a private research instrument, not a treatment product. It can collect quality-tagged sessions, record sham opportunities, deliver cues through a fail-closed protocol, and export traceable observations. Local verification supports software reliability only. It does not show deeper sleep or more lucid dreams.")
    heading(d,"Source mapping")
    rows=[
        ("Besedovsky 2017","Two 50 ms pink-noise tones, 1.075 s apart; EEG-triggered in consolidated NonREM; ACTIVE/SHAM crossover; n=14 healthy men.","Deep timing; sham design; conservative early-night eligibility.","No EEG. Timing preserved; phase locking absent."),
        ("Capstone details","One page; specific user/problem; numeric checks; Must/Won't; workflow; risk; human gate; three weeks.","Proposal, matrix, timeline, report and review gate.","Course guidance, not sleep evidence."),
        ("Android guidance","Persistent Data Items, transient Messages, capabilities, foreground-service limits.","Durable state/summary, direct live cue, user-started services.","Hardware/OS validation remains open."),
    ]
    t=d.add_table(rows=1,cols=4); t.style="Table Grid"
    fixed_table(t,[1600,3100,2800,usable_twips(d)-7590])
    for i,h in enumerate(("Source","Supported fact","Project use","Limit")):
        shade(t.rows[0].cells[i],NAVY); r=t.rows[0].cells[i].paragraphs[0].add_run(h); r.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,value in enumerate(row): shade(cells[i],"F5F8FA" if ri%2 else "FFFFFF"); margins(cells[i]); cells[i].paragraphs[0].add_run(value)
    fixed_table(t,[1600,3100,2800,usable_twips(d)-7590])
    heading(d,"Implemented evidence")
    bullets(d,["Versioned, bound, expiring, deduplicated cue protocol with explicit lifecycle.","Room on both devices for sessions, events, surveys, terminal results, and recovery.","Speaker/audio checks, preloaded deterministic waveforms, calibration, completion verification.","Signal quality, sleep-onset estimate, eligibility states, and heart-required lucid cues.","Blinded ACTIVE/SHAM, conditioning, morning report, history/export/delete, replay fingerprint."])
    heading(d,"Failed and incomplete sessions")
    d.add_paragraph("No physical overnight session was run in this implementation pass, so there is no honest failed-night count. The app records stopped, timed out, phone lost, sensor degraded/failed, low battery, interrupted, and internal/storage failure. Hardware acceptance is PENDING, not PASS.")
    heading(d,"Limitations")
    bullets(d,["Motion and optical heart rate cannot identify EEG phase or clinically score sleep.","The source study was a small lab sample of healthy young men with EEG and in-ear tones; home speaker delivery differs.","The lucid feature is exploratory and not directly validated by the supplied deep-sleep paper.","Debug APKs/tests do not establish efficacy, long-term safety, battery behavior, or public readiness.","Codex can summarize de-identified exports, not validate sensor truth or replace review."])
    heading(d,"Recommended next experiment")
    kv(d,[("QUESTION","Do ACTIVE and SHAM nights differ in restedness, awakenings, cue awareness, recall, lucid result, or adverse effects?"),("DESIGN","12–20 nights after three accelerated runs: sensor-only baseline, then blinded randomized ACTIVE/SHAM; reveal after survey."),("PRIMARY DATA","Usable-session rate, cue lifecycle, signal coverage, awakenings, restedness, adverse effects. Lucidity exploratory."),("STOP RULES","Disable cues on phone/sensor/storage/battery/audio failure or repeated cue-linked arousal; stop study for meaningful adverse effects."),("ANALYSIS","Paired descriptive differences with every failed/excluded session. No causal or medical claims.")])
    heading(d,"Human review checklist")
    bullets(d,["Verify source statements against DOI 10.1038/s41467-017-02170-3.","Confirm exports are de-identified before Codex use.","Attach hardware results and failed-session reasons.","Approve/reject each conclusion; sign and date."])
    d.add_paragraph("Reviewer: ____________________   Decision: ☐ Approved  ☐ Revise   Date: ____________")
    heading(d,"References")
    bullets(d,["Besedovsky L, Ngo H-VV, Dimitrov S, et al. Nature Communications. 2017;8:1984. doi:10.1038/s41467-017-02170-3.","Android Developers: Wear Data Layer client types; Discover devices; Foreground-service launch restrictions. Accessed 18 July 2026.","Dream Sync local research, capstone, code, schemas, and test outputs."])
    d.save(DOCS / "Dream Sync Codex Research Report.docx")

if __name__ == "__main__":
    DOCS.mkdir(exist_ok=True); proposal(); matrix(); report(); print("Generated 3 capstone documents")
